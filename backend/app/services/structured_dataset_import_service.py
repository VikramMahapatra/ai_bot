"""Synchronous bounded inspection; no file persistence, formulas or macros execution."""
import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree as ET
from fastapi import HTTPException
from app.schemas.structured_dataset import MAX_ROWS, MAX_COLUMNS, MAX_FILE_BYTES

MAX_TEXT = 4000
MAX_EXPANDED = 20 * 1024 * 1024
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def fail(message: str) -> None:
    raise HTTPException(422, message)


def infer(values: list) -> str:
    values = [v for v in values if v is not None and v != ""]
    if not values:
        return "text"
    if all(isinstance(v, bool) for v in values):
        return "boolean"
    if all(isinstance(v, (date, datetime)) for v in values):
        return "date"
    strings = [str(v) for v in values]
    # Text identifiers with zero padding must never become numeric.
    if any(re.match(r"^[+-]?0\d", v) for v in strings):
        return "text"
    if all(re.fullmatch(r"[+-]?\d+", v) for v in strings):
        return "integer"
    if all(re.fullmatch(r"[+-]?\d+(?:\.\d+)?", v) for v in strings):
        return "decimal"
    return "text"


def table(sheet_name: str, matrix: list[list], header_row: int) -> dict:
    if len(sheet_name) > 200:
        fail("Sheet name exceeds 200 characters")
    if len(matrix) < header_row:
        fail(f"{sheet_name}: header row does not exist")
    headers = matrix[header_row - 1]
    labels = [str(v).strip() if v is not None else "" for v in headers]
    if not labels or any(not v for v in labels):
        fail(f"{sheet_name}: headers must be nonempty")
    if len({v.casefold() for v in labels}) != len(labels):
        fail(f"{sheet_name}: headers must be unique")
    if len(labels) > MAX_COLUMNS or any(len(v) > 200 for v in labels):
        fail(f"{sheet_name}: header limits exceeded")
    data = matrix[header_row:]
    if len(data) > MAX_ROWS:
        fail(f"{sheet_name}: row limit exceeded")
    if any(len(row) > len(labels) for row in data):
        fail(f"{sheet_name}: row contains cells without a header")
    keys = []
    for label in labels:
        key = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")[:55]
        if not key or not key[0].isalpha():
            key = "column_" + key
        base, n = key, 2
        while key in keys:
            key = f"{base[:63 - len(str(n))]}_{n}"
            n += 1
        keys.append(key)
    columns = []
    for i, (key, label) in enumerate(zip(keys, labels)):
        columns.append(dict(key=key, label=label, data_type=infer([r[i] for r in data if len(r) > i]),
                            aliases=[], public=False, queryable=False))
    rows = []
    for raw in data:
        row = {}
        for i, key in enumerate(keys):
            value = raw[i] if i < len(raw) else None
            if value == "" and columns[i]["data_type"] != "text":
                value = None
            if isinstance(value, (date, datetime)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            elif isinstance(value, Decimal):
                value = format(value, "f")
            if value is not None and columns[i]["data_type"] == "text" and not isinstance(value, str):
                value = str(value)
            if value is not None and len(str(value)) > MAX_TEXT:
                fail(f"{sheet_name}: cell text exceeds {MAX_TEXT} characters")
            row[key] = value
        rows.append(row)
    return dict(sheet_name=sheet_name, columns=columns, rows=rows, row_count=len(rows))


def inspect_file(content: bytes, filename: str, header_row: int = 1,
                 sheet_name: str | None = None, discovery_only: bool = False) -> dict:
    if not 1 <= header_row <= 100:
        fail("header_row must be between 1 and 100")
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "File exceeds 5 MB")
    if sheet_name is not None and (not sheet_name or len(sheet_name) > 200):
        fail("sheet_name must contain 1–200 characters")
    extension = filename.lower().rsplit(".", 1)[-1]
    warnings = ["Approval required: all columns default to private and nonqueryable."]
    if extension == "csv":
        try:
            text = content.decode("utf-8-sig")
            matrix = []
            for row in csv.reader(io.StringIO(text), strict=True):
                if len(row) > MAX_COLUMNS or len(matrix) >= MAX_ROWS + header_row:
                    fail("CSV exceeds row/column limits")
                if any(len(v) > MAX_TEXT for v in row):
                    fail("CSV cell text limit exceeded")
                matrix.append(row)
            if sheet_name is not None and sheet_name != "CSV":
                fail("Requested sheet does not exist: " + sheet_name)
            return dict(sheet_names=["CSV"],
                        tables=[] if discovery_only else [table("CSV", matrix, header_row)],
                        warnings=[] if discovery_only else warnings)
        except (UnicodeError, csv.Error) as exc:
            fail(f"Invalid UTF-8 CSV: {exc}")
    if extension != "xlsx":
        fail("Only .xlsx and UTF-8 .csv are supported")
    try:
        # Inspect ZIP/XML before openpyxl allocates workbook structures.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 1000 or sum(i.file_size for i in infos) > MAX_EXPANDED:
                fail("Workbook expanded size limit exceeded")
            if len({i.filename for i in infos}) != len(infos):
                fail("Duplicate workbook ZIP entries")
            numeric = {}
            total_cells = 0
            for info in infos:
                if "vba" in info.filename.lower() or "externallinks/" in info.filename.lower():
                    fail("Macros and external workbook links are unsupported")
                if not info.filename.endswith(".xml"):
                    continue
                xml = archive.read(info)
                markup = xml.upper().replace(b"\x00", b"")  # also UTF-16/32 markup
                if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
                    fail("XML entities are unsupported")
                if not re.fullmatch(r"xl/worksheets/[^/]+\.xml", info.filename):
                    continue
                root = ET.fromstring(xml)
                for xml_row in root.findall(".//s:sheetData/s:row", NS):
                    row_ref = xml_row.get("r", "")
                    if not row_ref.isdigit() or not 1 <= int(row_ref) <= MAX_ROWS + header_row:
                        fail("Workbook actual row coordinates exceed limits")
                cells = root.findall(".//s:c", NS)
                total_cells += len(cells)
                if total_cells > MAX_ROWS * MAX_COLUMNS + 5000:
                    fail("Workbook overall cell limit exceeded")
                values = {}
                seen = set()
                for cell in cells:
                    coordinate = cell.get("r", "")
                    if coordinate in seen:
                        fail("Duplicate worksheet cell reference")
                    seen.add(coordinate)
                    match = re.fullmatch(r"([A-Z]+)([1-9]\d*)", coordinate)
                    if match is None:
                        fail("Invalid worksheet cell reference")
                    col = 0
                    for letter in match[1]:
                        col = col * 26 + ord(letter) - ord("A") + 1
                    if col > MAX_COLUMNS or int(match[2]) > MAX_ROWS + header_row:
                        fail("Workbook actual cell coordinates exceed limits")
                    value = cell.find("s:v", NS)
                    if cell.find("s:f", NS) is not None:
                        if value is None or value.text is None or cell.get("t") == "e":
                            fail("Formula has no usable cached result; convert formulas to values")
                        warnings.append("Cached formula results used without recalculation; verify freshness.")
                    if cell.get("t", "n") == "n" and value is not None and value.text:
                        number = Decimal(value.text)
                        # Bound BEFORE fixed-point rendering or openpyxl's float
                        # conversion. Tiny ZIP/XML values can have huge exponents.
                        if (not number.is_finite() or number.adjusted() > 99
                                or number.as_tuple().exponent < -100
                                or len(number.as_tuple().digits) > 200):
                            fail("Workbook numeric precision exceeds 100 integer/fractional digits")
                        values[coordinate] = number
                numeric[info.filename] = values
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            tables = []
            total_rows = 0
            if len(workbook.worksheets) > 20:
                fail("Workbook sheet limit exceeded")
            sheet_names = [sheet.title for sheet in workbook.worksheets]
            if any(len(name) > 200 for name in sheet_names):
                fail("Sheet name exceeds 200 characters")
            if sheet_name is not None and sheet_name not in sheet_names:
                fail("Requested sheet does not exist: " + sheet_name)
            for sheet in workbook.worksheets:
                # Ignore attacker-controlled dimension hints; actual cells were
                # preflighted above. Reset avoids silently dropping understated cells.
                sheet.reset_dimensions()
                matrix = []
                row_count = 0
                inspect_sheet = not discovery_only and (sheet_name is None or sheet.title == sheet_name)
                part = sheet._worksheet_path.lstrip("/")
                if part not in numeric:
                    fail("Unsupported worksheet part path")
                raw_numbers = numeric[part]
                for row in sheet.iter_rows():
                    if len(row) > MAX_COLUMNS or row_count >= MAX_ROWS + header_row:
                        fail(f"{sheet.title}: worksheet dimensions exceed limits")
                    row_count += 1
                    values = []
                    for cell in row:
                        value = cell.value
                        if cell.data_type == "e":
                            fail(f"{sheet.title}: Excel error cell")
                        if getattr(cell, "coordinate", None) in raw_numbers and not isinstance(value, (date, datetime)):
                            value = raw_numbers[cell.coordinate]
                            # Preserve Excel zero-padding identifier formats.
                            if re.fullmatch(r"0{2,}", cell.number_format) and value == value.to_integral_value():
                                value = str(int(value)).zfill(len(cell.number_format))
                        if value is not None and len(str(value)) > MAX_TEXT:
                            fail(f"{sheet.title}: cell text exceeds {MAX_TEXT} characters")
                        values.append(value)
                    if inspect_sheet:
                        matrix.append(values)
                # Every sheet counts toward resource bounds, even when it is
                # not selected. Only selected headers are interpreted as a table.
                total_rows += max(0, row_count - header_row)
                if total_rows > MAX_ROWS:
                    fail("Workbook overall row limit exceeded; no sheets were truncated")
                if not inspect_sheet:
                    continue
                width = max((len(r) for r in matrix), default=0)
                # Streaming rows can vary in width after reset_dimensions.
                matrix = [r + [None] * (width - len(r)) for r in matrix]
                result = table(sheet.title, matrix, header_row)
                tables.append(result)
            return dict(sheet_names=sheet_names, tables=tables,
                        warnings=[] if discovery_only else list(dict.fromkeys(warnings)))
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (ValueError, TypeError, OverflowError, InvalidOperation, KeyError, zipfile.BadZipFile, ET.ParseError, OSError) as exc:
        fail(f"Invalid XLSX: {exc}")