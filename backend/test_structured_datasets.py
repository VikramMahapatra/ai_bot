"""Run with unittest. All writes use SQLite :memory:; never import app.main.

Router loaded directly to avoid app.api's eager unrelated LLM/service imports.
PostgreSQL is validated by DDL compilation only, not a configured connection.
Migration entry points are not executed, even on SQLite.
"""
import asyncio
import importlib.util
import io
import pathlib
import sys
import types
import unittest
import zipfile
from unittest.mock import patch, Mock
import httpx
from fastapi import FastAPI, HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, event, insert, inspect as sa_inspect
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from sqlalchemy.dialects import postgresql
from sqlalchemy.schema import CreateTable, AddConstraint, CreateIndex
from sqlalchemy import create_mock_engine
from sqlalchemy.sql.ddl import SchemaGenerator
from sqlalchemy.exc import IntegrityError
from app.config import settings
from app.database import Base, get_db
from app.models import Organization, User, WidgetConfig, QualificationTemplate, UserRole
from app.models.structured_dataset import StructuredDataset as Dataset, StructuredDatasetVersion as Version, StructuredDatasetRow as Row
from app.models.structured_dataset import owned_metadata
from app.schemas.structured_dataset import DatasetCreate, VersionCreate, QueryPlan
# app.services.__init__ initializes unrelated persistent Chroma state. Load only
# owned service modules inside a temporary package shell, with no legacy edits.
services_package = types.ModuleType("app.services")
services_package.__path__ = [str(pathlib.Path(__file__).parent / "app/services")]
with patch.dict(sys.modules, {"app.services": services_package}):
    from app.services import structured_dataset_service as service
    from app.services.structured_dataset_import_service import inspect_file


def payload(widget="agent-a", amount="0.10", required=True):
    return DatasetCreate(widget_id=widget, name="Prices", description="Exact rates", sheet_name="Rates",
        columns=[dict(key="code", label="Code", data_type="text", public=True, queryable=True),
                 dict(key="price", label="Price", data_type="decimal", currency="USD", public=True, queryable=True),
                 dict(key="secret", label="Secret", data_type="text", public=False, queryable=False)],
        rows=[dict(code="001", price=amount, secret="hidden"), dict(code="002", price="0.20", secret="internal")],
        required_filters=["code"] if required else [])


def xlsx(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    data = io.BytesIO()
    workbook.save(data)
    workbook.close()
    return data.getvalue()


def rewrite_xlsx(source, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(output, "w") as new:
        for item in old.infolist():
            new.writestr(item, transform(item.filename, old.read(item)))
    return output.getvalue()


def multisheet_xlsx(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    data = io.BytesIO()
    workbook.save(data)
    workbook.close()
    return data.getvalue()


class ImportTests(unittest.TestCase):
    def test_discovery_and_selected_differing_headers(self):
        source = multisheet_xlsx([
            ("Coworking", [["Title"], ["Notes"], [], ["City", "Low", "High"], ["Ahmedabad", 5000, 9000]]),
            ("Private", [["Title"], [], ["Location", "Carpet rent", "Saleable rent"],
                         ["Andheri East", "₹180–400", "₹110–240"], ["Andheri West", "₹180–350", "₹110–210"]]),
        ])
        with patch.dict(sys.modules, {"app.services": services_package}), patch(
                "app.services.structured_dataset_import_service.table", side_effect=AssertionError("No table parsing during discovery")):
            self.assertEqual(inspect_file(source, "flexo.xlsx", discovery_only=True),
                             dict(sheet_names=["Coworking", "Private"], tables=[], warnings=[]))
        coworking = inspect_file(source, "flexo.xlsx", 4, "Coworking")
        self.assertEqual(coworking["sheet_names"], ["Coworking", "Private"])
        self.assertEqual(len(coworking["tables"]), 1)
        self.assertEqual(coworking["tables"][0]["rows"], [{"city": "Ahmedabad", "low": "5000", "high": "9000"}])
        private = inspect_file(source, "flexo.xlsx", 3, "Private")["tables"][0]
        self.assertEqual(private["row_count"], 2)
        self.assertEqual(private["rows"][0], {"location": "Andheri East", "carpet_rent": "₹180–400", "saleable_rent": "₹110–240"})
        self.assertTrue(all(not c["public"] and not c["queryable"] for c in private["columns"]))
        # Compatibility: omitted selection still applies the explicit row to all sheets.
        self.assertEqual(len(inspect_file(source, "flexo.xlsx", 4)["tables"]), 2)
        with self.assertRaisesRegex(HTTPException, "nonempty"):
            inspect_file(source, "flexo.xlsx", 3)
        with self.assertRaisesRegex(HTTPException, "header row does not exist"):
            inspect_file(source, "flexo.xlsx", 100, "Private")
        for name in ["Missing", "private", "", "x" * 201]:
            with self.subTest(name=name), self.assertRaises(HTTPException):
                inspect_file(source, "flexo.xlsx", 3, name)

    def test_selected_malformed_headers_and_data_not_skipped(self):
        for rows, error in [([["a", "A"], [1, 2]], "unique"),
                            ([["a", None], [1, 2]], "nonempty"),
                            ([["a"], [1, 2]], "nonempty")]:
            source = multisheet_xlsx([("Good", [["a"], [1]]), ("Bad", rows)])
            self.assertEqual(inspect_file(source, "a.xlsx", sheet_name="Good")["tables"][0]["row_count"], 1)
            with self.subTest(error=error), self.assertRaisesRegex(HTTPException, error):
                inspect_file(source, "a.xlsx", sheet_name="Bad")

    def test_csv_discovery_selection_and_legacy_calls(self):
        content = b"Title\ncode,price\n001,0.10\n"
        self.assertEqual(inspect_file(content, "a.csv", discovery_only=True),
                         dict(sheet_names=["CSV"], tables=[], warnings=[]))
        self.assertEqual(inspect_file(content, "a.csv", 2, "CSV"), inspect_file(content, "a.csv", 2))
        for discovery in [False, True]:
            with self.assertRaisesRegex(HTTPException, "does not exist"):
                inspect_file(content, "a.csv", 2, "Other", discovery)
            for invalid in [b"a\n\xff", b"a\n" + b"x\n" * 2001, b"a\n" + b"x" * 4001, b"," * 50]:
                with self.subTest(discovery=discovery), self.assertRaises(HTTPException):
                    inspect_file(invalid, "a.csv", discovery_only=discovery)

    def test_unselected_security_and_bounds_in_discovery(self):
        source = multisheet_xlsx([("Good", [["a"], [1]]), ("Other", [["a"], [1]])])
        attacks = [
            (lambda data: data.replace(b"<v>1</v>", b"<v>1e4001</v>"), "precision"),
            (lambda data: data.replace(b"<v>1</v>", b"<v>1e-4001</v>"), "precision"),
            (lambda data: data.replace(b"<v>1</v>", b"<v>NaN</v>"), "precision"),
            (lambda data: data.replace(b"<v>1</v>", b"<f>1+2</f><v></v>"), "cached"),
            (lambda data: data.replace(b't="n"><v>1</v>', b't="e"><v>#VALUE!</v>'), "error cell"),
            (lambda data: data.replace(b"</row></sheetData>", b'<c r="A2" t="n"><v>2</v></c></row></sheetData>'), "Duplicate"),
            (lambda data: data.replace(b'r="A2"', b'r="AY2"'), "coordinates"),
            (lambda data: data.replace(b'r="A2"', b'r="A2002"'), "coordinates"),
            (lambda data: data.replace(b'r="2"', b'r="2002"'), "coordinates"),
            (lambda data: data.replace(b"<t>a</t>", b"<t>" + b"x" * 4001 + b"</t>"), "cell text"),
            (lambda data: b'<!DOCTYPE worksheet [<!ENTITY x "abc">]>' + data, "entities"),
        ]
        for transform, error in attacks:
            data = rewrite_xlsx(source, lambda name, data: transform(data) if name == "xl/worksheets/sheet2.xml" else data)
            for options in [dict(sheet_name="Good"), dict(discovery_only=True)]:
                with self.subTest(error=error, options=options), self.assertRaisesRegex(HTTPException, error):
                    inspect_file(data, "a.xlsx", **options)
        for path in ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"]:
            output = io.BytesIO(source)
            with zipfile.ZipFile(output, "a") as archive:
                archive.writestr(path, b"inert")
            for options in [dict(sheet_name="Good"), dict(discovery_only=True)]:
                with self.subTest(path=path, options=options), self.assertRaisesRegex(HTTPException, "Macros"):
                    inspect_file(output.getvalue(), "a.xlsx", **options)
        too_many = multisheet_xlsx([(f"Sheet{n}", [["a"], [1]]) for n in range(21)])
        for options in [dict(sheet_name="Sheet0"), dict(discovery_only=True)]:
            with self.assertRaisesRegex(HTTPException, "sheet limit"):
                inspect_file(too_many, "a.xlsx", **options)

    def test_discovery_and_selection_cannot_bypass_archive_limits(self):
        source = xlsx([["a"], [1]])
        from app.schemas.structured_dataset import MAX_FILE_BYTES
        for options in [dict(sheet_name="Sheet"), dict(discovery_only=True)]:
            for content, filename in [(b"x" * (MAX_FILE_BYTES + 1), "a.xlsx"),
                                      (b"bad zip", "a.xlsx"), (source, "a.xlsm")]:
                with self.subTest(filename=filename, options=options), self.assertRaises(HTTPException):
                    inspect_file(content, filename, **options)
            for extra in ["expanded", "entries", "duplicate"]:
                output = io.BytesIO(source)
                with zipfile.ZipFile(output, "a", compression=zipfile.ZIP_DEFLATED) as archive:
                    if extra == "expanded":
                        archive.writestr("large.bin", b"x" * (20 * 1024 * 1024))
                    elif extra == "entries":
                        for n in range(1001):
                            archive.writestr(f"entry{n}.bin", b"")
                    else:
                        # Deliberately malformed ZIP; no filesystem persistence.
                        with patch("warnings.warn"):
                            archive.writestr("xl/workbook.xml", b"duplicate")
                with self.subTest(extra=extra, options=options), self.assertRaises(HTTPException):
                    inspect_file(output.getvalue(), "a.xlsx", **options)

    def test_nonstandard_sheet_part_retains_precision(self):
        source = xlsx([["price"], [1]])
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(output, "w") as new:
            for item in old.infolist():
                name, data = item.filename, old.read(item)
                if name == "xl/worksheets/sheet1.xml":
                    name = "xl/worksheets/rates.xml"
                    data = data.replace(b"<v>1</v>", b"<v>1.123456789123456789</v>")
                elif name == "xl/_rels/workbook.xml.rels":
                    data = data.replace(b"sheet1.xml", b"rates.xml")
                new.writestr(name, data)
        self.assertEqual(inspect_file(output.getvalue(), "a.xlsx")["tables"][0]["rows"][0]["price"], "1.123456789123456789")

    def test_utf16_entity_declarations_rejected(self):
        data = rewrite_xlsx(xlsx([["a"], ["x"]]), lambda name, data:
            ('<?xml version="1.0" encoding="utf-16"?><!DOCTYPE worksheet [<!ENTITY x "abc">]>'
             + data.decode()).encode("utf-16") if name == "xl/worksheets/sheet1.xml" else data)
        with self.assertRaisesRegex(HTTPException, "entities"):
            inspect_file(data, "a.xlsx")

    def test_blank_numeric_and_mixed_text_round_trip(self):
        result = inspect_file(b"code,price\n001,0.10\n002,\n", "a.csv")["tables"][0]
        self.assertIsNone(result["rows"][1]["price"])
        mixed = inspect_file(xlsx([["mixed"], [True], ["label"], [12]]), "a.xlsx")["tables"][0]
        self.assertTrue(all(isinstance(r["mixed"], str) for r in mixed["rows"]))
        for table in [result, mixed]:
            for column in table["columns"]:
                column["public"] = True
            _, report = service.validate_payload(VersionCreate.model_validate(
                {k: table[k] for k in ["sheet_name", "columns", "rows"]}))
            self.assertTrue(report["valid"], report)

    def test_xlsx_nonfinite_and_exponent_bounds(self):
        source = xlsx([["price"], [1]])
        for token in [b"NaN", b"Infinity", b"1e4001", b"1e-4001"]:
            data = rewrite_xlsx(source, lambda name, data: data.replace(b"<v>1</v>", b"<v>" + token + b"</v>")
                                if name == "xl/worksheets/sheet1.xml" else data)
            with self.subTest(token=token), self.assertRaises(HTTPException):
                inspect_file(data, "a.xlsx")

    def test_xlsx_duplicate_coordinates(self):
        data = rewrite_xlsx(xlsx([["price"], [1]]), lambda name, data:
            data.replace(b"</row></sheetData>", b'<c r="A2" t="n"><v>2</v></c></row></sheetData>')
            if name == "xl/worksheets/sheet1.xml" else data)
        with self.assertRaises(HTTPException):
            inspect_file(data, "a.xlsx")

    def test_csv_decimal_and_leading_zeros(self):
        result = inspect_file(b"code,price\r\n001,123456789.123456789\r\n", "rates.csv")
        table = result["tables"][0]
        self.assertEqual(table["rows"][0], {"code": "001", "price": "123456789.123456789"})
        self.assertEqual([c["data_type"] for c in table["columns"]], ["text", "decimal"])
        self.assertTrue(all(not c["public"] and not c["queryable"] for c in table["columns"]))

    def test_csv_headers_bounds_and_utf8(self):
        for content in [b"a,A\n1,2", b"a,\n1,2", b"a\n1,2", b"a\n\xff", b"a\n" + b"x\n" * 2001]:
            with self.subTest(content=content[:20]), self.assertRaises(HTTPException):
                inspect_file(content, "a.csv")

    def test_csv_formula_is_inert_text(self):
        self.assertEqual(inspect_file(b"formula\n=SUM(A1:A2)", "a.csv")["tables"][0]["rows"][0]["formula"], "=SUM(A1:A2)")

    def test_xlsx_exact_decimal_and_zero_format(self):
        workbook = Workbook()
        workbook.active.append(["Code", "Price"])
        workbook.active.append([1, 0.1])
        workbook.active["A2"].number_format = "000"
        data = io.BytesIO()
        workbook.save(data)
        workbook.close()
        result = inspect_file(data.getvalue(), "prices.xlsx")["tables"][0]
        self.assertEqual(result["rows"][0], {"code": "001", "price": "0.1"})

    def test_xlsx_formula_missing_cache(self):
        with self.assertRaisesRegex(HTTPException, "cached"):
            inspect_file(xlsx([["price"], ["=1+2"]]), "a.xlsx")

    def test_cached_formula_and_exact_xml_numeric(self):
        source = xlsx([["price"], ["=1+2"]])
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(output, "w") as new:
            for item in old.infolist():
                data = old.read(item)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<v></v>", b"<v>123456789.123456789123</v>")
                new.writestr(item, data)
        result = inspect_file(output.getvalue(), "a.xlsx")
        self.assertEqual(result["tables"][0]["rows"][0]["price"], "123456789.123456789123")
        self.assertTrue(any("Cached formula" in w for w in result["warnings"]))

    def test_workbook_overall_limit_no_truncation(self):
        workbook = Workbook()
        workbook.active.append(["a"])
        for n in range(1001):
            workbook.active.append([n])
        second = workbook.create_sheet("Second")
        second.append(["a"])
        for n in range(1000):
            second.append([n])
        data = io.BytesIO()
        workbook.save(data)
        workbook.close()
        for options in [{}, dict(sheet_name=workbook.worksheets[0].title), dict(discovery_only=True)]:
            with self.subTest(options=options), self.assertRaisesRegex(HTTPException, "overall row"):
                inspect_file(data.getvalue(), "a.xlsx", **options)

    def test_sparse_sheet_and_bad_archives(self):
        self.assertEqual(inspect_file(xlsx([["a", "b"], ["x", None]]), "a.xlsx")["tables"][0]["rows"][0]["b"], None)
        for data, name in [(b"not zip", "bad.xlsx"), (b"a", "bad.xlsm"), (b"x" * (5 * 1024 * 1024 + 1), "big.csv")]:
            with self.assertRaises(HTTPException):
                inspect_file(data, name)

    def test_understated_dimensions_and_macro_content(self):
        source = xlsx([["a"], ["one"], ["two"]])
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(output, "w") as new:
            for item in old.infolist():
                data = old.read(item)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'ref="A1:A3"', b'ref="A1:A1"')
                new.writestr(item, data)
        self.assertEqual(inspect_file(output.getvalue(), "a.xlsx")["tables"][0]["row_count"], 2)
        with zipfile.ZipFile(output, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"inert")
        with self.assertRaisesRegex(HTTPException, "Macros"):
            inspect_file(output.getvalue(), "a.xlsx")


class DatasetFixture(unittest.TestCase):
    def setUp(self):
        self.flags = patch.multiple(settings, STRUCTURED_DATASETS_ENABLED=True,
            STRUCTURED_DATASETS_CHAT_ENABLED=True, STRUCTURED_DATASETS_ORGANIZATION_IDS="1, 2")
        self.flags.start()
        self.engine = create_engine("sqlite:///:memory:", poolclass=StaticPool,
                                    connect_args={"check_same_thread": False})
        @event.listens_for(self.engine, "connect")
        def foreign_keys(connection, record):
            connection.execute("PRAGMA foreign_keys=ON")
        Base.metadata.create_all(self.engine, tables=[Organization.__table__, User.__table__,
            QualificationTemplate.__table__, WidgetConfig.__table__])
        with self.engine.begin() as conn:
            owned = [Dataset.__table__, Version.__table__, Row.__table__]
            Base.metadata.create_all(conn, tables=owned, checkfirst=True)
            Base.metadata.create_all(conn, tables=owned, checkfirst=True)
            conn.execute(insert(Organization), [dict(id=1, name="One", org_domain="one"), dict(id=2, name="Two", org_domain="two")])
            conn.execute(insert(User), [dict(id=1, organization_id=1, username="one", email="one", hashed_password="unused", role=UserRole.ADMIN),
                                      dict(id=2, organization_id=2, username="two", email="two", hashed_password="unused", role=UserRole.ADMIN)])
            conn.execute(insert(QualificationTemplate), [dict(id=1, organization_id=1, name="one", objective="test"),
                                                         dict(id=2, organization_id=2, name="two", objective="test")])
            conn.execute(insert(WidgetConfig), [dict(user_id=1, organization_id=1, widget_id="agent-a", name="A", qualification_template_id=1),
                dict(user_id=1, organization_id=1, widget_id="agent-b", name="B", qualification_template_id=1),
                dict(user_id=2, organization_id=2, widget_id="agent-other", name="Other", qualification_template_id=2)])
        self.db = Session(self.engine, expire_on_commit=False)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()
        self.flags.stop()

    def published(self, data=None):
        dataset, version = service.create_dataset(self.db, 1, data or payload())
        self.db.commit()
        service.publish_version(self.db, 1, dataset.id, version.id)
        self.db.commit()
        return dataset, version

    def query(self, dataset, **kwargs):
        plan = dict(dataset_id=dataset.id, operation="lookup", filters=[dict(column="code", operator="eq", value="001")])
        plan.update(kwargs)
        return service.execute_query(self.db, 1, "agent-a", plan)


class DatabaseTests(DatasetFixture):
    def test_scoped_child_and_active_foreign_keys(self):
        dataset, version = self.published()
        other, other_version = service.create_dataset(self.db, 2, payload(widget="agent-other"))
        self.db.commit()
        with self.assertRaises(IntegrityError):
            self.db.execute(insert(Row), dict(dataset_id=dataset.id, version_id=version.id,
                organization_id=2, widget_id="agent-other", row_number=99, values={}))
        self.db.rollback()
        dataset.active_version_id = other_version.id
        with self.assertRaises(IntegrityError):
            self.db.flush()
        self.db.rollback()

    def test_range_types_and_empty_numeric_aggregate(self):
        data = payload(required=False).model_dump()
        data["columns"].append(dict(key="unused", label="Unused", data_type="decimal", public=True, queryable=True))
        data = DatasetCreate.model_validate(data)
        dataset, _ = self.published(data)
        with self.assertRaises(HTTPException):
            self.query(dataset, filters=[dict(column="code", operator="gte", value="001")])
        result = self.query(dataset, operation="sum", filters=[], aggregate_column="unused")
        self.assertEqual(result["status"], "no_match")
        self.assertEqual(result["rows"], [])

    def test_published_inserts_and_deletes_guarded(self):
        dataset, version = self.published()
        self.db.add(Row(dataset_id=dataset.id, version_id=version.id, organization_id=1,
                        widget_id="agent-a", row_number=99, values={}))
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()
        self.db.delete(service.rows_query(self.db, version).first())
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()
        self.db.delete(version)
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()

    def test_expired_published_version_cannot_be_demoted(self):
        _, version = self.published()
        self.db.expire(version, ["status"])
        version.status = "draft"
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()

    def test_expired_published_row_cannot_move_to_draft(self):
        dataset, version = self.published()
        draft = service.add_version(self.db, dataset, VersionCreate.model_validate(
            payload().model_dump(include={"sheet_name", "columns", "rows", "required_filters"})))
        self.db.commit()
        row = service.rows_query(self.db, version).first()
        self.db.expire(row, ["version_id"])
        row.version_id, row.row_number = draft.id, 99
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()

    def test_large_integer_json_round_trip_and_numeric_comparison(self):
        data = payload(required=False)
        data.columns[1].data_type = "integer"
        data.rows[0]["price"], data.rows[1]["price"] = "9007199254740993", "9007199254740994"
        dataset, version = self.published(data)
        self.assertEqual(service.serialize_dataset(self.db, dataset)["versions"][0]["rows"][0]["price"], "9007199254740993")
        self.assertTrue(service.validate_version(self.db, version)["valid"])
        result = self.query(dataset, filters=[dict(column="price", operator="gte", value="9007199254740994")])
        self.assertEqual(result["rows"][0]["price"], "9007199254740994")
        self.assertEqual(self.query(dataset, operation="sum", aggregate_column="price", filters=[])["rows"][0]["value"], "18014398509481987")

    def test_required_filters_need_nonblank_equality(self):
        dataset, _ = self.published()
        # A range on a required numeric field must not unlock an entire table.
        second_data = payload()
        second_data.required_filters = ["price"]
        second, _ = self.published(second_data)
        for target, filters in [(dataset, [dict(column="code", operator="eq", value=" ")]),
                                (second, [dict(column="price", operator="gte", value="0")])]:
            self.assertEqual(self.query(target, filters=filters)["status"], "clarify")

    def test_no_match_aggregates_do_not_fabricate_zero(self):
        dataset, _ = self.published()
        for operation in ["count", "sum", "min", "max"]:
            kw = dict(operation=operation, filters=[dict(column="code", operator="eq", value="absent")])
            if operation != "count":
                kw["aggregate_column"] = "price"
            result = self.query(dataset, **kw)
            self.assertEqual((result["status"], result["rows"], result["citations"]), ("no_match", [], []))

    def test_request_string_bounds_and_pg_json_strings(self):
        for change in [dict(select=["x" * 65]), dict(aggregate_column="x" * 65),
                       dict(filters=[dict(column="", operator="eq", value="a")])]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                QueryPlan.model_validate(dict(dataset_id=1, operation="lookup", **change))
        for field, value in [("required_filters", ["x" * 65]), ("name", "bad\x00name"),
                             ("description", "bad\ud800")]:
            with self.subTest(field=field), self.assertRaises(ValueError):
                DatasetCreate.model_validate(dict(payload().model_dump(), **{field: value}))
        data = payload().model_dump()
        data["columns"][0]["aliases"] = ["x" * 201]
        with self.assertRaises(ValueError):
            DatasetCreate.model_validate(data)
        data = payload().model_dump()
        data["rows"][0]["price"] = 9007199254740993
        with self.assertRaises(ValueError):
            DatasetCreate.model_validate(data)
        with self.assertRaises(ValueError):
            QueryPlan.model_validate(dict(dataset_id=1, operation="count", filters=[
                dict(column="price", operator="eq", value=9007199254740993)]))

    def test_publish_lookup_catalog(self):
        dataset, version = self.published()
        result = self.query(dataset)
        self.assertEqual(result["rows"], [{"code": "001", "price": "0.10"}])
        self.assertIn("USD", result["response"])
        self.assertEqual(result["citations"][0]["row_numbers"], [1])
        self.assertNotIn("secret", str(service.get_active_catalog(self.db, 1, "agent-a")))

    def test_exact_sum_ranges_count_min_max(self):
        dataset, _ = self.published(payload(required=False))
        for operation, expected in [("sum", "0.30"), ("min", "0.10"), ("max", "0.20")]:
            self.assertEqual(self.query(dataset, operation=operation, filters=[], aggregate_column="price")["rows"][0]["value"], expected)
        self.assertEqual(self.query(dataset, operation="count", filters=[dict(column="price", operator="gte", value="0.20")])["rows"], [{"count": 1}])

    def test_private_fields_blocked_in_all_positions(self):
        dataset, _ = self.published()
        for kw in [dict(select=["secret"]), dict(filters=[dict(column="secret", operator="eq", value="hidden")]),
                   dict(operation="sum", aggregate_column="secret"), dict(select=["forged"])]:
            with self.subTest(kw=kw), self.assertRaises(HTTPException):
                self.query(dataset, **kw)

    def test_public_nonqueryable_filter_and_aggregate_blocked(self):
        data = payload(required=False)
        data.columns[1].queryable = False
        dataset, _ = self.published(data)
        self.assertIn("price", self.query(dataset, filters=[])["rows"][0])
        for kw in [dict(filters=[dict(column="price", operator="eq", value="0.10")]),
                   dict(operation="sum", filters=[], aggregate_column="price")]:
            with self.assertRaises(HTTPException):
                self.query(dataset, **kw)

    def test_publish_checks_actual_data_not_cached_report(self):
        dataset, version = service.create_dataset(self.db, 1, payload())
        self.db.commit()
        row = service.rows_query(self.db, version).first()
        row.values = dict(row.values, price="corrupt")
        self.db.commit()
        self.assertTrue(version.validation["valid"])
        with self.assertRaises(HTTPException):
            service.publish_version(self.db, 1, dataset.id, version.id)
        self.db.rollback()
        self.assertIsNone(dataset.active_version_id)

    def test_quotas_bound_history_and_tenant_storage(self):
        dataset, _ = self.published()
        data = VersionCreate.model_validate(payload().model_dump(include={"sheet_name", "columns", "rows", "required_filters"}))
        with patch.object(service, "MAX_VERSIONS", 1), self.assertRaises(HTTPException):
            service.add_version(self.db, dataset, data)
        with patch.object(service, "MAX_TENANT_ROWS", 2), self.assertRaises(HTTPException):
            service.add_version(self.db, dataset, data)
        with patch.object(service, "MAX_TENANT_DATASETS", 1), self.assertRaises(HTTPException):
            service.create_dataset(self.db, 1, payload())

    def test_required_filters_no_match_and_ambiguity(self):
        dataset, _ = self.published()
        self.assertEqual(self.query(dataset, filters=[])["status"], "clarify")
        self.assertEqual(self.query(dataset, filters=[dict(column="code", operator="eq", value="missing")])["status"], "no_match")
        result = self.query(dataset, limit=1, filters=[dict(column="code", operator="eq", value="001")])
        self.assertEqual(result["total_matches"], 1)
        second, _ = self.published(payload(required=False))
        result = self.query(second, filters=[], limit=1)
        self.assertIn("Showing 1 of 2", result["response"])

    def test_tenant_agent_and_forged_version_isolation(self):
        dataset, version = self.published()
        for call in [lambda: service.get_dataset(self.db, 2, dataset.id),
                     lambda: service.execute_query(self.db, 1, "agent-b", dict(dataset_id=dataset.id, operation="count")),
                     lambda: service.create_dataset(self.db, 1, payload(widget="agent-other")),
                     lambda: service.get_version(self.db, dataset, version.id + 99)]:
            with self.assertRaises(HTTPException) as error:
                call()
            self.assertEqual(error.exception.status_code, 404)

    def test_drafts_old_live_and_published_immutability(self):
        dataset, old = self.published()
        service.lock_tenant(self.db, 1)
        new = service.add_version(self.db, dataset, VersionCreate.model_validate(payload(amount="9.99").model_dump(include={"sheet_name", "columns", "rows", "required_filters"})))
        self.db.commit()
        self.assertEqual(self.query(dataset)["rows"][0]["price"], "0.10")
        service.publish_version(self.db, 1, dataset.id, new.id)
        self.db.commit()
        self.assertEqual(self.query(dataset)["rows"][0]["price"], "9.99")
        self.assertEqual(old.status, "published")
        with self.assertRaises(HTTPException):
            self.query(dataset, version_id=old.id)
        with self.assertRaises(HTTPException):
            service.publish_version(self.db, 1, dataset.id, old.id)
        self.db.rollback()
        old.sheet_name = "changed"
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()
        row = service.rows_query(self.db, old).first()
        row.values = {"code": "changed"}
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.flush()
        self.db.rollback()

    def test_raw_invalid_draft_and_publish_revalidation(self):
        data = payload(amount="not a number")
        dataset, version = service.create_dataset(self.db, 1, data)
        self.db.commit()
        self.assertFalse(version.validation["valid"])
        self.assertEqual(service.rows_query(self.db, version).first().values["price"], "not a number")
        with self.assertRaises(HTTPException) as error:
            service.publish_version(self.db, 1, dataset.id, version.id)
        self.assertEqual(error.exception.status_code, 422)
        self.assertIsNone(dataset.active_version_id)

    def test_validation_duplicate_keys_types_unknown_and_required(self):
        data = payload()
        data.columns.append(data.columns[0])
        data.rows[0]["unknown"] = "x"
        data.required_filters = ["secret"]
        _, report = service.validate_payload(data)
        self.assertFalse(report["valid"])
        self.assertTrue(any("Duplicate column keys" in e for e in report["errors"]))
        self.assertTrue(any("unknown column" in e for e in report["errors"]))
        for value, kind in [("2026-99-01", "date"), ("yes", "boolean"), ("1.5", "integer"), (1.5, "decimal")]:
            with self.assertRaises(ValueError):
                service.normalize(value, kind)
        self.assertEqual(service.normalize("true", "boolean"), True)
        self.assertEqual(service.normalize("2026-10-13", "date"), "2026-10-13")

    def test_disabled_no_database_queries(self):
        db = Mock()
        with patch.object(settings, "STRUCTURED_DATASETS_ENABLED", False):
            self.assertEqual(service.get_active_catalog(db, 1, "agent-a"), [])
            with self.assertRaises(HTTPException):
                service.execute_query(db, 1, "agent-a", {})
            db.query.assert_not_called()
        with patch.object(settings, "STRUCTURED_DATASETS_ORGANIZATION_IDS", ""):
            self.assertFalse(service.is_structured_enabled(1))
        with patch.object(settings, "STRUCTURED_DATASETS_CHAT_ENABLED", False):
            self.assertEqual(service.get_active_catalog(db, 1, "agent-a"), [])
            db.query.assert_not_called()

    def test_soft_delete_inaccessible(self):
        dataset, _ = self.published()
        dataset.is_deleted = True
        self.db.commit()
        self.assertEqual(service.get_active_catalog(self.db, 1, "agent-a"), [])
        with self.assertRaises(HTTPException):
            self.query(dataset)

    def test_strict_query_schema(self):
        for kw in [dict(sql="SELECT *"), dict(limit=21), dict(limit=True), dict(dataset_id="1"),
                   dict(filters=[dict(column="price", operator="eq", value=0.1)])]:
            plan = dict(dataset_id=1, operation="count")
            plan.update(kw)
            with self.assertRaises(ValueError):
                QueryPlan.model_validate(plan)

    def test_postgres_ddl_and_owned_migration(self):
        names = set(sa_inspect(self.engine).get_table_names())
        self.assertEqual(names, {"organizations", "users", "qualification_templates", "widget_configs",
                                 "structured_datasets", "structured_dataset_versions", "structured_dataset_rows"})
        dialect = postgresql.dialect()
        for table in [Dataset.__table__, Version.__table__, Row.__table__]:
            ddl = str(CreateTable(table).compile(dialect=dialect))
            if table is Dataset.__table__:
                self.assertIn("id INTEGER GENERATED BY DEFAULT AS IDENTITY", ddl)
            else:
                self.assertIn("SERIAL", ddl)
                self.assertIn("JSONB", ddl)
            for index in table.indexes:
                self.assertIn("CREATE INDEX", str(CreateIndex(index).compile(dialect=dialect)))
        # AddConstraint mutates its constraint's inline-create rule. Compile a
        # DDL-only copy, never the shared model used by subsequent SQLite tests.
        constraint = next(c for c in owned_metadata().tables["structured_datasets"].constraints if c.name == "fk_sd_active_scope")
        self.assertIn("active_version_id", str(AddConstraint(constraint).compile(dialect=dialect)))

    def test_offline_pg_creation_orders_deferred_fk_and_repeat_is_empty(self):
        statements, existing = [], set()
        def capture(sql, *args, **kwargs):
            statements.append(str(sql.compile(dialect=postgresql.dialect())))
            if isinstance(sql, CreateTable):
                existing.add(sql.element.name)
        engine = create_mock_engine("postgresql://", capture)
        with patch.object(engine.dialect, "has_table", side_effect=lambda conn, name, **kw: name in existing):
            # MockConnection.create_all forces checkfirst=False; exercise the real
            # DDL visitor with a fake catalog rather than misreporting idempotency.
            metadata = owned_metadata()
            owned = [metadata.tables[t.__tablename__] for t in [Dataset, Version, Row]]
            SchemaGenerator(engine.dialect, engine, checkfirst=True, tables=owned).traverse_single(metadata)
            first = len(statements)
            SchemaGenerator(engine.dialect, engine, checkfirst=True, tables=owned).traverse_single(metadata)
        self.assertEqual(len(statements), first)
        self.assertEqual(sum("CREATE TABLE" in sql for sql in statements), 3)
        self.assertFalse(any("CREATE TYPE" in sql for sql in statements))
        self.assertEqual(sum("ADD CONSTRAINT fk_sd_active_scope" in sql for sql in statements), 1)
        fk_at = next(i for i, sql in enumerate(statements) if "ADD CONSTRAINT fk_sd_active_scope" in sql)
        self.assertTrue(all(i < fk_at for i, sql in enumerate(statements) if "CREATE TABLE" in sql))


class ApiTests(DatasetFixture):
    def test_inspect_discovery_selected_sheet_contract_and_scope(self):
        path = pathlib.Path(__file__).parent / "app/api/structured_datasets.py"
        spec = importlib.util.spec_from_file_location("isolated_inspect_router", path)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"app.services": services_package}):
            spec.loader.exec_module(module)
        app = FastAPI()
        app.include_router(module.router)
        app.dependency_overrides[module.require_admin] = lambda: self.db.get(User, 1)
        app.dependency_overrides[get_db] = lambda: self.db
        source = multisheet_xlsx([
            ("Coworking", [["Title"], ["Notes"], [], ["City", "Price"], ["Ahmedabad", "5000"]]),
            ("Private", [["Title"], [], ["Location", "Rent"], ["Andheri East", "₹180–400"]]),
        ])

        async def requests():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                async def inspect(**fields):
                    return await client.post("/api/admin/knowledge/datasets/inspect",
                        data={"widget_id": "agent-a", **fields}, files={"file": ("flexo.xlsx", source)})
                with event_query_recorder(self.engine) as queries:
                    discovery = await inspect(discovery_only="true")
                    self.assertEqual(discovery.status_code, 200, discovery.text)
                    self.assertEqual(discovery.json(), dict(sheet_names=["Coworking", "Private"], tables=[], warnings=[]))
                    for sheet, header in [("Coworking", "4"), ("Private", "3")]:
                        response = await inspect(sheet_name=sheet, header_row=header)
                        self.assertEqual(response.status_code, 200, response.text)
                        self.assertEqual(response.json()["sheet_names"], ["Coworking", "Private"])
                        self.assertEqual([t["sheet_name"] for t in response.json()["tables"]], [sheet])
                    self.assertEqual(response.json()["tables"][0]["rows"][0], {"location": "Andheri East", "rent": "₹180–400"})
                    self.assertEqual((await inspect(header_row="4")).status_code, 200)
                    self.assertEqual((await inspect(sheet_name="Missing", header_row="3")).status_code, 422)
                    self.assertEqual((await inspect(sheet_name="Private", header_row="100")).status_code, 422)
                    self.assertEqual((await inspect(header_row="0", discovery_only="true")).status_code, 422)
                    self.assertEqual((await inspect(discovery_only="not-a-bool")).status_code, 422)
                    csv = await client.post("/api/admin/knowledge/datasets/inspect",
                        data={"widget_id": "agent-a", "discovery_only": "true"}, files={"file": ("a.csv", b"Title\na,A\n1,2")})
                    self.assertEqual(csv.json(), dict(sheet_names=["CSV"], tables=[], warnings=[]))
                    csv = await client.post("/api/admin/knowledge/datasets/inspect",
                        data={"widget_id": "agent-a", "sheet_name": "Other"}, files={"file": ("a.csv", b"a\n1")})
                    self.assertEqual(csv.status_code, 422)
                    self.assertEqual((await inspect(widget_id="agent-other", discovery_only="true")).status_code, 404)
                    with patch.object(settings, "STRUCTURED_DATASETS_ENABLED", False):
                        self.assertEqual((await inspect(discovery_only="true")).status_code, 403)
                self.assertFalse(any("INSERT " in q.upper() or "UPDATE " in q.upper() or "DELETE " in q.upper() for q in queries))
                self.assertEqual(self.db.query(Dataset).count(), 0)
        asyncio.run(requests())

    def test_api_lifecycle_and_disabled(self):
        path = pathlib.Path(__file__).parent / "app/api/structured_datasets.py"
        spec = importlib.util.spec_from_file_location("isolated_dataset_router", path)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"app.services": services_package}):
            spec.loader.exec_module(module)
        app = FastAPI()
        app.include_router(module.router)
        user = self.db.get(User, 1)
        app.dependency_overrides[module.require_admin] = lambda: user
        app.dependency_overrides[get_db] = lambda: self.db
        prefix = "/api/admin/knowledge/datasets"

        async def requests():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                with event_query_recorder(self.engine) as queries:
                    overlarge = await client.post(prefix + "/", content=b" " * (5 * 1024 * 1024 + 1), headers={"content-type": "application/json"})
                    self.assertEqual(overlarge.status_code, 413)
                    async def chunks():
                        for _ in range(6):
                            yield b" " * (1024 * 1024)
                    overlarge = await client.post(prefix + "/", content=chunks(), headers={"content-type": "application/json", "content-length": "1"})
                    self.assertEqual(overlarge.status_code, 413)
                self.assertEqual(queries, [])
                self.assertEqual((await client.get(prefix + "/capabilities?widget_id=agent-a")).json()["enabled"], True)
                inspected = await client.post(prefix + "/inspect", data={"widget_id": "agent-a"}, files={"file": ("a.csv", b"code,price\n001,0.10")})
                self.assertEqual(inspected.status_code, 200, inspected.text)
                created = await client.post(prefix + "/", json=payload().model_dump())
                self.assertEqual(created.status_code, 201, created.text)
                data = created.json()
                self.assertEqual(data["status"], "draft")
                self.assertEqual(data["draft_version"]["status"], "draft")
                self.assertEqual(data["draft_version"]["rows"][0]["price"], "0.10")
                id, version_id = data["id"], data["draft_version"]["id"]
                self.assertEqual((await client.patch(f"{prefix}/{id}", json={"status": "active"})).status_code, 409)
                self.assertTrue((await client.post(f"{prefix}/{id}/versions/{version_id}/validate")).json()["valid"])
                published = await client.post(f"{prefix}/{id}/versions/{version_id}/publish")
                self.assertEqual(published.status_code, 200)
                self.assertEqual(published.json()["status"], "active")
                self.assertEqual(published.json()["versions"][0]["status"], "published")
                draft_payload = payload(amount="1.23").model_dump(include={"sheet_name", "columns", "rows", "required_filters"})
                draft_response = await client.post(f"{prefix}/{id}/versions", json=draft_payload)
                self.assertEqual(draft_response.status_code, 201, draft_response.text)
                self.assertEqual(draft_response.json()["active_version_id"], version_id)
                self.assertEqual(len((await client.get(f"{prefix}/{id}")).json()["versions"]), 2)
                self.assertEqual((await client.patch(f"{prefix}/{id}", json={"status": "inactive"})).status_code, 200)
                inactive = await client.post(f"{prefix}/{id}/test-query", json=dict(dataset_id=id, operation="count"))
                self.assertEqual(inactive.status_code, 409)
                patched = await client.patch(f"{prefix}/{id}", json={"status": "active", "name": "New name"})
                self.assertEqual(patched.status_code, 200)
                self.assertEqual(len(patched.json()["versions"]), 2)
                self.assertEqual((await client.patch(f"{prefix}/{id}", json={"status": None})).status_code, 422)
                self.assertEqual((await client.get(f"{prefix}/{id}/rows?skip=1&limit=1")).json()["items"][0]["row_number"], 2)
                result = await client.post(f"{prefix}/{id}/test-query", json=dict(dataset_id=id, operation="lookup", filters=[]))
                self.assertEqual(result.json()["status"], "clarify")
                self.assertEqual((await client.get(prefix + "/?widget_id=agent-b")).json()["items"], [])
                other = self.db.get(User, 2)
                app.dependency_overrides[module.require_admin] = lambda: other
                self.assertEqual((await client.get(f"{prefix}/{id}")).status_code, 404)
                app.dependency_overrides[module.require_admin] = lambda: user
                self.assertEqual((await client.delete(f"{prefix}/{id}")).status_code, 200)
                self.assertEqual((await client.get(f"{prefix}/{id}")).status_code, 404)
                with patch.object(settings, "STRUCTURED_DATASETS_ENABLED", False):
                    with event_query_recorder(self.engine) as queries:
                        response = await client.get(prefix + "/capabilities")
                        self.assertFalse(response.json()["enabled"])
                        self.assertFalse(response.json()["chat_enabled"])
                        self.assertEqual((await client.get(prefix + "/?widget_id=agent-a")).json()["items"], [])
                        self.assertEqual((await client.post(prefix + "/", json=payload().model_dump())).status_code, 403)
                    self.assertFalse(any("structured_dataset" in q for q in queries))
        asyncio.run(requests())

    def test_real_admin_dependency_denies_nonadmin(self):
        from app.auth import require_admin
        with self.assertRaises(HTTPException) as error:
            require_admin(User(role=UserRole.USER))
        self.assertEqual(error.exception.status_code, 403)


class event_query_recorder:
    def __init__(self, engine):
        self.engine, self.queries = engine, []
    def record(self, conn, cursor, statement, parameters, context, many):
        self.queries.append(statement)
    def __enter__(self):
        event.listen(self.engine, "before_cursor_execute", self.record)
        return self.queries
    def __exit__(self, *args):
        event.remove(self.engine, "before_cursor_execute", self.record)


if __name__ == "__main__":
    unittest.main()